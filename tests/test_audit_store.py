from __future__ import annotations
import io
import zipfile
from pathlib import Path
import pytest
import openpyxl
import audit_store as audit


def base_row(name="Product", url="https://example.com/p"):
    return {
        "main_category":"Home Appliances", "subcategory":"Ovens", "page_label":"Page 1",
        "page_url":"https://example.com/cat", "product_name":name, "product_url":url,
        "status":"review", "operator_decision":"", "consistency_score":"80",
        "problem_summary":"Needs work", "suggested_fix":"Fix it", "profile":"boxy", "subtype":"boxy",
        "current_image_file":"", "normalized_image_file":"", "comparison_image_file":"",
    }

@pytest.fixture
def audit_paths(tmp_path, monkeypatch):
    root = tmp_path / "audit_data"
    monkeypatch.setattr(audit, "AUDIT_DIR", root)
    monkeypatch.setattr(audit, "MASTER_AUDIT_PATH", root / "master_audit.csv")
    return root


def test_same_name_different_urls_media_do_not_overwrite(audit_paths):
    a = base_row("Same Name", "https://example.com/a")
    b = base_row("Same Name", "https://example.com/b")
    paths_a = audit.save_audit_media(a, b"A", b"AN", b"AC")
    paths_b = audit.save_audit_media(b, b"B", b"BN", b"BC")
    assert paths_a["current_image_file"] != paths_b["current_image_file"]
    assert (audit_paths / paths_a["current_image_file"]).read_bytes() == b"A"
    assert (audit_paths / paths_b["current_image_file"]).read_bytes() == b"B"


def test_same_name_same_url_media_is_predictable(audit_paths):
    a = base_row("Same Name", "https://example.com/a")
    b = base_row("Same Name", "https://example.com/a")
    assert audit.media_directory_for_row(a) == audit.media_directory_for_row(b)

@pytest.mark.parametrize("name", ["", "Café 冰箱", "CON", "NUL", "COM1", "LPT1", "x"*300])
def test_storage_names_are_safe(audit_paths, name):
    path = audit.media_directory_for_row(base_row(name, "https://example.com/" + str(abs(hash(name)))))
    assert len(path.name) <= 120
    assert path.name.rstrip(" .") == path.name
    assert not path.name.upper().split("__",1)[0] in {"CON","NUL","COM1","LPT1"}


def test_duplicate_batch_rows_are_written_once(audit_paths):
    row = base_row()
    result = audit.update_master_rows([row, dict(row)], mode="append")
    assert result["added"] == 1
    assert len(audit.load_master_rows()) == 1


def test_duplicate_across_append_operations_is_not_readded(audit_paths):
    row = base_row()
    audit.update_master_rows([row], mode="append")
    result = audit.update_master_rows([dict(row)], mode="append")
    assert result["added"] == 0
    assert len(audit.load_master_rows()) == 1


def test_same_name_different_urls_remain_separate(audit_paths):
    audit.update_master_rows([base_row("Same", "https://example.com/a"), base_row("Same", "https://example.com/b")], mode="append")
    assert len(audit.load_master_rows()) == 2


def test_url_normalization_dedupes_equivalent_urls(audit_paths):
    a = base_row("X", "https://EXAMPLE.com:443/p/")
    b = base_row("X", "https://example.com/p")
    audit.update_master_rows([a,b], mode="append")
    assert len(audit.load_master_rows()) == 1

@pytest.mark.parametrize("payload", ["=1+1", "+SUM(A1:A2)", "-CMD|' /C calc'!A0", "@something", "normal text", "https://example.com/x"])
def test_excel_untrusted_strings_are_not_formulas(audit_paths, payload):
    row = base_row(payload, "https://example.com/p")
    data = audit.build_boss_excel_report([row])
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = wb["Simple Audit"]
    headers = {cell.value: cell.column for cell in ws[2]}
    cell = ws.cell(row=3, column=headers["Product Name"])
    assert cell.value == payload
    assert cell.data_type != "f"


def test_numeric_values_remain_numbers_in_safe_writer(audit_paths):
    # Directly test an intentionally numeric field by using build report source row.
    row = base_row(); row["component_count"] = 3
    data = audit.build_boss_excel_report([row])
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = wb["Advanced Audit"]
    headers = {cell.value: cell.column for cell in ws[2]}
    cell = ws.cell(row=3, column=headers["Component Count"])
    assert cell.value == 3
    assert cell.data_type == "n"


def test_zip_does_not_include_traversal_file(audit_paths, tmp_path):
    secret = tmp_path / "secret.txt"; secret.write_text("SECRET")
    row = base_row(); row["current_image_file"] = "../secret.txt"
    package = audit.build_master_visual_package([row])
    with zipfile.ZipFile(io.BytesIO(package)) as zf:
        names = zf.namelist()
        assert not any("secret.txt" in n.lower() for n in names)
        assert "UNSAFE_OR_MISSING_MEDIA_SKIPPED.txt" in names


def test_zip_rejects_absolute_external_path(audit_paths, tmp_path):
    secret = tmp_path / "outside.png"; secret.write_bytes(b"x")
    row = base_row(); row["current_image_file"] = str(secret)
    package = audit.build_master_visual_package([row])
    with zipfile.ZipFile(io.BytesIO(package)) as zf:
        assert not any(n.endswith("outside.png") for n in zf.namelist())


def test_zip_includes_valid_nested_media(audit_paths):
    nested = audit_paths / "media" / "safe" / "x.png"
    nested.parent.mkdir(parents=True); nested.write_bytes(b"PNGDATA")
    row = base_row(); row["current_image_file"] = "media/safe/x.png"
    package = audit.build_master_visual_package([row])
    with zipfile.ZipFile(io.BytesIO(package)) as zf:
        assert "media/safe/x.png" in zf.namelist()
        assert zf.read("media/safe/x.png") == b"PNGDATA"


def test_zip_rejects_symlink_escape(audit_paths, tmp_path):
    outside = tmp_path / "outside.png"; outside.write_bytes(b"secret")
    link = audit_paths / "media" / "link.png"; link.parent.mkdir(parents=True)
    try: link.symlink_to(outside)
    except OSError: pytest.skip("symlink unavailable")
    row = base_row(); row["current_image_file"] = "media/link.png"
    package = audit.build_master_visual_package([row])
    with zipfile.ZipFile(io.BytesIO(package)) as zf:
        assert not any(n.endswith("link.png") for n in zf.namelist())
